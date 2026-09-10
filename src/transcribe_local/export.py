# -*- coding: utf-8 -*-
"""导出：md / txt / srt / docx。

`[❓]` 一律保留 —— 那是复核入口，为了「好看」去掉等于把存疑处埋了。
"""
from __future__ import annotations

import re
from pathlib import Path

LINE = re.compile(r"^\[(?P<a>[\d:.]+)\s*-\s*(?P<b>[\d:.]+)\]\s*(?P<spk>[^:]+):\s*(?P<text>.*)$")
CREDIT = "由 Transcribe Local 生成 · transcribe.solutions"


def parse(transcript: str) -> list[tuple[str, str, str, str]]:
    rows = []
    for ln in transcript.splitlines():
        m = LINE.match(ln.strip())
        if m:
            rows.append((m["a"], m["b"], m["spk"].strip(), m["text"].strip()))
    return rows


def _secs(t: str) -> float:
    mm, ss = t.split(":")
    return int(mm) * 60 + float(ss)


def _srt_time(x: float) -> str:
    h, rem = divmod(x, 3600)
    m, s = divmod(rem, 60)
    return f"{int(h):02d}:{int(m):02d}:{int(s):02d},{int(round((s - int(s)) * 1000)):03d}"


def write(transcript: str, out_base: Path, cfg: dict) -> list[Path]:
    rows = parse(transcript)
    credit = bool(cfg["export"].get("footer_credit", True))
    written: list[Path] = []
    for fmt in cfg["export"].get("formats", ["md"]):
        p = out_base.with_suffix("." + fmt)
        if fmt == "md":
            body = "\n\n".join(f"**{spk}** `{a}`\n\n{text}" for a, _, spk, text in rows)
            p.write_text(body + (f"\n\n---\n\n<sub>{CREDIT}</sub>\n" if credit else "\n"), encoding="utf-8")
        elif fmt == "txt":
            body = "\n".join(f"[{a}] {spk}: {text}" for a, _, spk, text in rows)
            p.write_text(body + (f"\n\n{CREDIT}\n" if credit else "\n"), encoding="utf-8")
        elif fmt == "srt":
            p.write_text("".join(
                f"{i}\n{_srt_time(_secs(a))} --> {_srt_time(_secs(b))}\n{spk}: {text}\n\n"
                for i, (a, b, spk, text) in enumerate(rows, 1)), encoding="utf-8")
        elif fmt == "docx":
            _docx(rows, p, credit)
        else:
            continue
        written.append(p)
    return written


def _docx(rows, path: Path, credit: bool) -> None:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError as e:
        raise RuntimeError("导出 .docx 需要 python-docx：pip install 'transcribe-local[docx]'") from e
    doc = Document()
    for a, _, spk, text in rows:
        para = doc.add_paragraph()
        head = para.add_run(f"{spk}  {a}\n")
        head.bold = True
        head.font.size = Pt(9)
        head.font.color.rgb = RGBColor(0x75, 0x65, 0x4F)
        para.add_run(text)
    if credit:
        run = doc.add_paragraph().add_run(CREDIT)
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(0xB5, 0xA8, 0x8F)
    doc.save(path)
