"""大纲文件读取（2026-08-19）。

守两件事：**顺序**与**拒绝**。
顺序——大纲里的表格常夹在小节中间，先取段落再取表格会让模型读到的上下文跟原件对不上；
拒绝——坏文件/空文件/超大文件必须给出人话理由，不能抛到 500 去。
"""
import io

import pytest

from app import outline_file


def _docx(builder) -> bytes:
    """现造一个 .docx（用导出时就已装好的 python-docx，测试不引新依赖）。"""
    from docx import Document
    doc = Document()
    builder(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_reads_paragraphs_and_tables_in_document_order():
    def build(doc):
        doc.add_paragraph("一、渠道现状")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "术语"
        t.cell(0, 1).text = "说明"
        t.cell(1, 0).text = "头水货"
        t.cell(1, 1).text = "首批到店的车源"
        doc.add_paragraph("二、人物")

    text = outline_file.read("大纲.docx", _docx(build))
    lines = text.split("\n")
    # 表格必须落在两个小节之间——先段落后表格的写法会把「二、人物」顶到表格前面
    assert lines[0] == "一、渠道现状"
    assert lines[-1] == "二、人物"
    assert any("头水货" in l and "首批到店的车源" in l for l in lines)


def test_drops_blank_paragraphs_and_blank_table_rows():
    def build(doc):
        doc.add_paragraph("有字")
        doc.add_paragraph("   ")
        doc.add_paragraph("")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "留下"          # 另一行整行空白
        doc.add_paragraph("结尾")

    text = outline_file.read("a.docx", _docx(build))
    assert text.split("\n") == ["有字", "留下", "结尾"]


def test_rejects_unsupported_extension():
    with pytest.raises(outline_file.OutlineReadError) as e:
        outline_file.read("大纲.pdf", b"%PDF-1.4")
    assert e.value.code == "outline_not_docx"


def test_rejects_a_file_that_is_not_really_docx():
    # 改个扩展名就想混进来的（常见：把 .doc 或 .pages 手动改名）
    with pytest.raises(outline_file.OutlineReadError):
        outline_file.read("假的.docx", b"this is plain text, not a zip")


# ── 接口层 ──────────────────────────────────────────────────────────────

def _client(monkeypatch):
    from fastapi.testclient import TestClient
    from app import api
    monkeypatch.setattr(api, "_current_email", lambda r: "u@example.com")
    return TestClient(api.app)


def test_endpoint_returns_text_and_char_count(monkeypatch):
    data = _docx(lambda d: d.add_paragraph("尚界与鸿蒙智行的分工"))
    with _client(monkeypatch) as c:
        r = c.post("/api/glossaries/outline",
                   files={"file": ("大纲.docx", data,
                                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["text"] == "尚界与鸿蒙智行的分工"
    assert body["chars"] == len(body["text"])
    assert body["truncated"] is False


def test_endpoint_truncates_at_the_draft_limit(monkeypatch):
    from app import glossary_assist
    over = glossary_assist.MAX_OUTLINE_CHARS + 500
    data = _docx(lambda d: d.add_paragraph("词" * over))
    with _client(monkeypatch) as c:
        r = c.post("/api/glossaries/outline", files={"file": ("长.docx", data, "application/octet-stream")})
    body = r.json()
    # 截断而不是拒绝：大纲前半通常最相关，直接拒了用户还得自己去裁
    assert body["truncated"] is True
    assert len(body["text"]) == glossary_assist.MAX_OUTLINE_CHARS


def test_endpoint_rejects_a_document_with_no_text(monkeypatch):
    data = _docx(lambda d: d.add_paragraph("   "))
    with _client(monkeypatch) as c:
        r = c.post("/api/glossaries/outline", files={"file": ("空.docx", data, "application/octet-stream")})
    assert r.status_code == 422
    assert r.json()["detail"]["code"] == "outline_no_text"


def test_endpoint_rejects_oversized_upload(monkeypatch):
    big = b"x" * (outline_file.MAX_BYTES + 1)
    with _client(monkeypatch) as c:
        r = c.post("/api/glossaries/outline", files={"file": ("大.docx", big, "application/octet-stream")})
    assert r.status_code == 413


def test_endpoint_requires_a_signed_in_user():
    """鉴权不能漏：这条路径会读用户上传的文件，必须跟其它术语库接口同一道闸。
    本机版没有登录，那道闸是「只接受本机访问」：外来的 Host 必须被拒。"""
    from fastapi.testclient import TestClient
    from app import api
    data = _docx(lambda d: d.add_paragraph("x"))
    with TestClient(api.app, base_url="http://attacker.example") as c:
        r = c.post("/api/glossaries/outline", files={"file": ("a.docx", data, "application/octet-stream")})
    assert r.status_code in (401, 403)
