"""转录结果导出：segments → Word docx；后处理产物 markdown → docx（简化段落渲染）。"""
import re
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from . import bidi, speaker_labels


def _set_font(run) -> None:
    run.font.name = "Cambria"                                # 英文/数字
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")    # 中文
    run.font.size = Pt(11)


def _set_rtl(para) -> None:
    """把一个段落设成从右往左。**两个元素都要，缺一不可**（OOXML 规范）：

    - 段落级 `<w:bidi/>`：管缩进 / 对齐 / 制表位，**它本身不改变段内文字的排列顺序**；
    - 文字级 `<w:rtl/>`：管这段文字自己的排列方向。

    只设 `w:bidi` 的症状是「对齐对了、字序还是错的」——看起来像修好了，其实没有。
    python-docx 没有现成接口，直接插 XML 元素；插的正是规范定义的那两个，不是绕路。
    """
    pPr = para._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)
    for run in para.runs:
        rPr = run._element.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rPr.append(rtl)


def segments_to_docx(segments: list[dict], lang: str = "", ui: str = "zh") -> bytes:
    """segments(每项含 t/s/sp) → docx 字节；说话人加粗，正文仿宋/Cambria 11pt。"""
    doc = Document()
    rtl = bidi.is_rtl(lang)
    for seg in segments:
        para = doc.add_paragraph()
        sp = speaker_labels.speaker_label(seg.get("sp"), ui)
        if sp:
            # ⚠️ 分隔符跟着界面语言走，别写死全角「：」——英文名后挂一个全角冒号
            # （`Interviewer：`）在任何语言下都是错的排版。txt 那边一直是对的，docx 漏了。
            r = para.add_run(f"{sp}{speaker_labels.speaker_sep(ui)}")
            r.font.bold = True
            _set_font(r)
        r = para.add_run(seg.get("s", ""))
        _set_font(r)
        if rtl:
            _set_rtl(para)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def markdown_to_docx(md_text: str, lang: str = "") -> bytes:
    """后处理产物（markdown）→ docx 的简化段落渲染：#/##/### 标题行去井号加粗，
    其余行原样成段，空行跳过。复用转录导出的字体设置（仿宋/Cambria 11pt）——
    不做完整 markdown 解析（产物是纪要/叙述稿，段落+标题已覆盖）。"""
    doc = Document()
    rtl = bidi.is_rtl(lang)
    for line in md_text.splitlines():
        s = line.strip()
        if not s:
            continue
        m = re.match(r"^(#{1,6})\s*(.*)$", s)
        para = doc.add_paragraph()
        run = para.add_run(m.group(2) if m else s)
        if m:
            run.font.bold = True
        _set_font(run)
        if rtl:
            _set_rtl(para)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
